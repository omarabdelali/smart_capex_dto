from datetime import date
import csv
from pathlib import Path
from tempfile import TemporaryDirectory
import pandas as pd
from django.test import SimpleTestCase,TestCase
from core.models import Site,Cell,MAOSMonthlyKPI,ConfigurationFeatures,RadioNetworkEvent,AnalysisEpisode,AnalysisEpisodeMember
from services.data_scope import plan_cband,explicit_cband,filter_cellplan_frame,filter_maos_chunk
from services.data_scope_cleanup import curate_maos,curate_plan,delete_site_data,inside_workspace
from services.models.prediction.analysis_episodes import refresh_episodes


POLICY={'enabled':True,'excluded_site_ids':['1','101'],'excluded_site_names':['REMOVE'],
        'excluded_cell_names':['cband','fdd_at_removed_site']}


class ScopeClassificationTests(SimpleTestCase):
    def test_cell_scope_retains_fdd_at_mixed_site_and_filters_tdd(self):
        policy={'enabled':True,'scope':'tdd_cells','excluded_cell_names':['td'], 'excluded_site_ids':[], 'excluded_site_names':[]}
        plan=pd.DataFrame([{'site_id':'1','cell_name':'td','band':'TDD'}, {'site_id':'1','cell_name':'fdd','band':'L1800'}])
        kept,stats=filter_cellplan_frame(plan,policy)
        self.assertEqual(list(kept.cell_name),['fdd'])
        self.assertEqual(stats['excluded_rows'],1)
        maos=pd.DataFrame({'Cell Name':['td','fdd','other_tdd'],'Frequency band':['42','3','40'],
            'Cell FDD TDD Indication':['CELL_TDD','CELL_FDD','CELL_TDD'],'eNodeB Name':['SITE_LM']*3})
        retained,count=filter_maos_chunk(maos,policy)
        self.assertEqual(list(retained['Cell Name']),['fdd'])
        self.assertEqual(count,2)

    def test_cband_is_not_every_tdd_band_or_an_arbitrary_frequency(self):
        self.assertTrue(explicit_cband('42.0'))
        self.assertTrue(explicit_cband('C-band'))
        self.assertTrue(plan_cband('TDD',43090,'4G'))
        self.assertTrue(plan_cband('TDD',45589,'LTE'))
        self.assertFalse(plan_cband('TDD',39600,'4G'))
        self.assertFalse(plan_cband('TDD',43090,'3G'))
        self.assertFalse(plan_cband('U2100',10588,'3G'))
        self.assertFalse(plan_cband('TDD',None))

    def test_entire_mixed_site_and_colocated_alias_are_filtered_before_loading(self):
        frame=pd.DataFrame([
            {'site_id':1,'site_name':'REMOVE','band':'TDD','frequency':43090,'technology':'LTE','latitude':35.,'longitude':9.},
            {'site_id':1,'site_name':'REMOVE','band':'L1800','frequency':1600,'technology':'LTE','latitude':35.,'longitude':9.},
            {'site_id':101,'site_name':'ALIAS','band':'U900','frequency':2959,'technology':'UMTS','latitude':35.00001,'longitude':9.},
            {'site_id':2,'site_name':'KEEP','band':'L1800','frequency':1600,'technology':'LTE','latitude':36.,'longitude':9.},
            {'site_id':3,'site_name':'OTHER_TDD','band':'TDD','frequency':39000,'technology':'LTE','latitude':34.,'longitude':9.},
        ])
        kept,stats=filter_cellplan_frame(frame,{'enabled':True})
        self.assertEqual(set(kept.site_id),{2,3})
        self.assertEqual(stats['excluded_rows'],3)

    def test_disabled_scope_does_not_change_existing_import_behavior(self):
        frame=pd.DataFrame({'site_id':[1]})
        kept,stats=filter_cellplan_frame(frame,{'enabled':False})
        self.assertIs(kept,frame)
        self.assertEqual(stats['excluded_rows'],0)

    def test_maos_excludes_other_bands_at_removed_site_and_keeps_nulls(self):
        frame=pd.DataFrame({'Cell Name':['cband','fdd_at_removed_site','unknown_fdd','keep'],
            'Frequency band':['42','3','1','3'],'eNodeB Name':['REMOVE_TD','REMOVE_LM','REMOVE_L2B','KEEP_LM'],
            'counter':[None,'0','1.23456789','0']})
        kept,count=filter_maos_chunk(frame,POLICY)
        self.assertEqual(count,3)
        self.assertEqual(kept.iloc[0]['Cell Name'],'keep')
        self.assertEqual(kept.iloc[0]['counter'],'0')

    def test_source_csv_curation_keeps_retained_counter_strings_exact(self):
        with TemporaryDirectory() as folder:
            source=Path(folder)/'source.csv';target=Path(folder)/'filtered.csv'
            with source.open('w',newline='',encoding='utf-8') as stream:
                stream.write('\n'*6)
                writer=csv.writer(stream)
                writer.writerow(['Cell Name',' Frequency band','eNodeB Name','counter'])
                writer.writerows([['fdd_at_removed_site','3','REMOVE_LM','999'],['keep','3','KEEP_LM','1.234567890123E+12'],['keep','3','KEEP_LM','']])
            report=curate_maos(source,target,POLICY)
            with target.open(newline='',encoding='utf-8') as stream:
                for _ in range(6):next(stream)
                rows=list(csv.DictReader(stream))
            self.assertEqual(report['rows_removed'],1)
            self.assertEqual(rows[0]['counter'],'1.234567890123E+12')
            self.assertEqual(rows[1]['counter'],'')

    def test_all_workbook_sheets_remove_site_rows(self):
        import openpyxl
        with TemporaryDirectory() as folder:
            source=Path(folder)/'plan.xlsx';target=Path(folder)/'filtered.xlsx'
            book=openpyxl.Workbook();sheet=book.active;sheet.title='Cells Info'
            sheet.append(['Site ID','Site Name','Cell Name','Band','Frequency','Technology'])
            sheet.append([1,'REMOVE','fdd_at_removed_site','L1800',1600,'LTE'])
            sheet.append([2,'KEEP','keep','L1800',1600,'LTE'])
            sheet=book.create_sheet('CellPlan');sheet.append(['SiteID','SiteName','CellName']);sheet.append([1,'REMOVE','fdd_at_removed_site']);sheet.append([2,'KEEP','keep'])
            book.save(source);curate_plan(source,target,POLICY)
            result=openpyxl.load_workbook(target,read_only=True)
            for sheet in result:
                rows=list(sheet.values);self.assertEqual(len(rows),2);self.assertEqual(rows[1][0],2)
            result.close()

    def test_mutation_paths_must_stay_in_workspace(self):
        with self.assertRaises(ValueError):inside_workspace(Path('/unrelated')/'file')


class ScopeCascadeTests(TestCase):
    def test_protected_episode_links_features_kpis_and_aliases_removed_without_touching_survivor(self):
        removed=[]
        for key,canonical in [('1','1'),('101','1'),('2','2')]:
            site=Site.objects.create(site_id=key,site_name=key,latitude=35.,longitude=9.,location='',canonical_site_id=canonical)
            cell=Cell.objects.create(site=site,cell_id=key,cell_name=key,technology='4G',frequency=1600,band='L1800',azimuth=0,antenna_height_m=25)
            MAOSMonthlyKPI.objects.create(cell=cell,month_date=date(2025,1,1),thrp_bits_dl=1.)
            ConfigurationFeatures.objects.create(site=site)
            RadioNetworkEvent.objects.create(site=site,event_key=key,event_type='CELL_ADDED',effective_date=date(2025,1,1),detector_version='fixture',details={'cell_ids':[key]})
            if key!='2':removed.append(site.pk)
        refresh_episodes();delete_site_data(removed)
        self.assertEqual(list(Site.objects.values_list('site_id',flat=True)),['2'])
        self.assertEqual(Cell.objects.count(),1)
        self.assertEqual(MAOSMonthlyKPI.objects.count(),1)
        self.assertEqual(ConfigurationFeatures.objects.count(),1)
        self.assertEqual(AnalysisEpisode.objects.count(),1)
        self.assertEqual(AnalysisEpisodeMember.objects.count(),1)
        self.assertEqual(RadioNetworkEvent.objects.count(),1)
