"""Audited whole-site scope removal with restorable inputs and derived-data rebuild."""
from datetime import datetime,timezone
import csv
import json
from pathlib import Path
import shutil
import sqlite3
from uuid import uuid4
from django.conf import settings
from django.db import transaction,connection
from django.db.models import F,Q
from django.core.cache import cache
from core.models import Site,Cell,MAOSMonthlyKPI,AnalysisEpisode,AnalysisEpisodeMember,FeatureSyncJob
from .data_scope import scope_policy,source_hash,explicit_cband,plan_cband,tdd_cell,normalized,filter_cellplan_frame,STANDARD_SOURCE


def inside_workspace(path):
    root=Path(settings.BASE_DIR).resolve();path=Path(path).resolve()
    if path==root or not path.is_relative_to(root):
        raise ValueError(f'Cleanup target is outside the project workspace: {path}')
    return path


def curate_maos(source,output,policy):
    """Preserve every retained counter string and the native six-line preamble."""
    ids=set(policy['excluded_cell_names']);names=set(policy['excluded_site_names'])
    before=after=0
    with Path(source).open(encoding='utf-8-sig',newline='') as src,Path(output).open('w',encoding='utf-8',newline='') as dst:
        for _ in range(6):
            dst.write(next(src))
        reader=csv.reader(src);writer=csv.writer(dst)
        header=next(reader);writer.writerow(header)
        columns={name.strip():i for i,name in enumerate(header)}
        for row in reader:
            before+=1
            value=lambda key:row[columns[key]].strip() if key in columns and columns[key]<len(row) else ''
            node=value('eNodeB Name')
            is_tdd=normalized(value('Cell FDD TDD Indication')) in ('TDD','CELLTDD')
            if value('Cell Name') in ids or explicit_cband(value('Frequency band')) or (policy.get('scope')=='tdd_cells' and is_tdd) or (policy.get('scope')!='tdd_cells' and (node in names or node.rsplit('_',1)[0] in names)):
                continue
            writer.writerow(row);after+=1
    return {'rows_before':before,'rows_retained':after,'rows_removed':before-after,'sha256':source_hash(output)}


def curate_plan(source,output,policy):
    import openpyxl
    original=openpyxl.load_workbook(source,read_only=True,data_only=True)
    result=openpyxl.Workbook(write_only=True);report={}
    ids=set(policy['excluded_site_ids']);names=set(policy['excluded_site_names']);cellnames=set(policy['excluded_cell_names'])
    for sheet in original:
        target=result.create_sheet(sheet.title);rows=sheet.iter_rows(values_only=True)
        header=next(rows);target.append(header)
        normalize=lambda text:''.join(c for c in str(text).upper() if c.isalnum())
        columns={normalize(name):i for i,name in enumerate(header)}
        value=lambda row,key:row[columns[key]] if key in columns else None
        kept=removed=0
        for row in rows:
            identifier=str(value(row,'SITEID')).removesuffix('.0')
            cell_scope=policy.get('scope')=='tdd_cells'
            if ((not cell_scope and (identifier in ids or str(value(row,'SITENAME')) in names)) or str(value(row,'CELLNAME')) in cellnames
                    or (tdd_cell(value(row,'BAND'),value(row,'FREQUENCY'),value(row,'TECHNOLOGY')) if cell_scope else plan_cband(value(row,'BAND'),value(row,'FREQUENCY'),value(row,'TECHNOLOGY')))):
                removed+=1;continue
            target.append(row);kept+=1
        report[sheet.title]={'rows_retained':kept,'rows_removed':removed}
    result.save(output);original.close()
    return {'sheets':report,'sha256':source_hash(output),'format':'retained values in all sheets; original formatting/formulas archived'}


def prepare_scope_cleanup(audit_path,progress=print,resume_backup=None,installed_plan_hash=None):
    report=json.loads(Path(audit_path).read_text(encoding='utf-8'))
    if report['scope']!='whole_physical_site':
        raise ValueError('Only the user-approved whole physical-site scope is supported.')
    if not report['site_rows_to_remove']:
        raise ValueError('No active C-band sites were identified; existing exclusions will not be reset.')
    if FeatureSyncJob.objects.filter(status__in=['queued','running']).exists():
        raise ValueError('A feature sync is active; finish it before cleanup.')
    root=Path(settings.BASE_DIR)
    backup=inside_workspace(root/resume_backup) if resume_backup else inside_workspace(root/'data/backups'/('cband_scope_'+datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S')+'_'+uuid4().hex[:6]))
    if not resume_backup:
        backup.mkdir(parents=True)
    policy={**report['policy'],'standard_source':STANDARD_SOURCE,'updated_at':datetime.now(timezone.utc).isoformat(),
        'backup':str(backup.relative_to(root)),'audit':str(Path(audit_path).resolve().relative_to(root.resolve()))}
    previous=scope_policy()
    for key in ('excluded_site_ids','excluded_site_names','excluded_cell_names'):
        policy[key]=sorted(set(policy[key])|set(previous.get(key,[])))
    import pandas as pd
    plan_path=(backup/'inputs'/report['cell_plan']['path']) if resume_backup else (root/report['cell_plan']['path'])
    frame=pd.read_excel(plan_path)
    frame.columns=frame.columns.str.strip().str.lower().str.replace(' ','_')
    _,source_scope=filter_cellplan_frame(frame,policy)
    blocked=set(source_scope['excluded_site_ids'])
    extra=Site.objects.filter(site_id__in=blocked).exclude(site_id__in=policy['excluded_site_ids'])
    if extra.exists():
        raise ValueError('The full source plan identifies additional active physical sites; expand the audit before removal.')
    policy['excluded_site_ids']=sorted(set(policy['excluded_site_ids'])|blocked)
    ids=frame['site_id'].astype(str).str.replace(r'\.0$','',regex=True)
    names=frame['site_name'].astype(str).str.strip()
    policy['excluded_cell_names']=sorted(set(policy['excluded_cell_names'])|set(frame.loc[ids.isin(blocked),'cell_name'].astype(str)))
    name_ids={name:set(ids[names==name]) for name in names.unique()}
    policy['excluded_site_names']=sorted(set(previous.get('excluded_site_names',[]))|{name for name,values in name_ids.items() if values and values<=set(policy['excluded_site_ids'])})
    policy['source_scope']=source_scope
    original_policy=root/'config/data_scope.json'
    if original_policy.exists() and not resume_backup:
        shutil.copy2(original_policy,backup/'previous_data_scope.json')
    if not resume_backup:
        shutil.copy2(audit_path,backup/'audit.json')
    db=Path(connection.settings_dict['NAME'])
    if 'sqlite' not in connection.settings_dict['ENGINE']:
        raise ValueError('This cleanup requires an explicitly prepared database backup for non-SQLite databases.')
    inside_workspace(db)
    if not resume_backup:
        with sqlite3.connect(str(db)) as src,sqlite3.connect(str(backup/'database.sqlite3')) as dst:
            src.backup(dst)
    progress(f'Consistent database backup: {backup}')
    sources=[report['cell_plan'],*report['source_reports']]
    prepared=[]
    for index,source in enumerate(sources):
        path=inside_workspace(root/source['path'])
        archived=backup/'inputs'/path.relative_to(root)
        temporary=inside_workspace(path.with_name(path.name+'.scope-filtered.tmp'))
        if resume_backup:
            if not archived.exists() or source_hash(archived)!=source['sha256']:
                raise ValueError(f'Backup does not match the audit: {archived}')
            if path.suffix.lower()=='.xlsx' and not temporary.exists():
                if not installed_plan_hash or source_hash(path)!=installed_plan_hash:
                    raise ValueError('Installed cell plan differs from the verified prepared copy.')
                info={'sha256':source_hash(path),'already_installed':True}
            else:
                if source_hash(path)!=source['sha256'] or not temporary.exists():
                    raise ValueError(f'Cannot safely resume source: {path}')
                # Recreate the staged file from the verified original to ensure
                # no partial or modified temporary output can be installed.
                progress(f'Verifying and rebuilding staged source: {path.name}')
                info=curate_plan(archived,temporary,policy) if path.suffix.lower()=='.xlsx' else curate_maos(archived,temporary,policy)
        else:
            if source_hash(path)!=source['sha256']:
                raise ValueError(f'Source changed since audit: {path}. Re-run the audit before removal.')
            archived.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(path,archived)
            progress(f'Preparing filtered source: {path.name}')
            info=curate_plan(path,temporary,policy) if path.suffix.lower()=='.xlsx' else curate_maos(path,temporary,policy)
        prepared.append((path,temporary,info))
    # All outputs/backup files exist before replacing active inputs. Database
    # removal remains atomic and starts only after source preparation succeeds.
    for path,temporary,info in prepared:
        if not info.get('already_installed'):
            temporary.replace(path)
        progress(f'Filtered input installed: {path.name}; {info}')
    (root/'config/data_scope.json').write_text(json.dumps(policy,indent=2),encoding='utf-8')
    current=set(Site.objects.filter(site_id__in=policy['excluded_site_ids']).values_list('pk',flat=True))
    if len(current)!=report['site_rows_to_remove']:
        raise ValueError('Site inventory changed since audit. Reconcile the audit before database removal.')
    count,details,deleted_members,deleted_episodes=delete_site_data(current)
    # Saved training/range/calibration outputs and disposable caches contain old
    # site identities or old neighborhood values. Preserve them only as backups.
    calibration=inside_workspace(root/'data/calibration');archived=backup/'calibration';archived.mkdir()
    archived_files=[]
    for path in list(calibration.iterdir()):
        if path.is_file():
            path=inside_workspace(path);shutil.copy2(path,archived/path.name);path.unlink();archived_files.append(path.name)
    cache.clear()
    from services.models.similarity.profile_cache import cache_backend
    backend=cache_backend();inside_workspace(backend._dir);backend.clear()
    result={'version':'exclude-cband-physical-sites-v1','backup':str(backup.relative_to(root)),
        'physical_sites_removed':report['physical_sites_to_remove'],'site_rows_removed':len(current),
        'deleted_records':details,'deleted_episode_members':deleted_members,'deleted_episodes':deleted_episodes,
        'total_cascade_records':count,'sources':[{**info,'path':str(path.relative_to(root))} for path,_,info in prepared],
        'invalidated_artifacts':archived_files,'remaining_sites':Site.objects.count(),'remaining_cells':Cell.objects.count(),
        'remaining_kpis':MAOSMonthlyKPI.objects.count(),'rebuild_status':'pending'}
    (backup/'removal.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    return result


def delete_site_data(site_pks):
    """Explicitly remove protected episode links before normal site cascades."""
    with transaction.atomic():
        episodes=AnalysisEpisode.objects.filter(Q(site_id__in=site_pks)|Q(raw_events__site_id__in=site_pks)).distinct()
        episode_ids=list(episodes.values_list('pk',flat=True))
        deleted_members=AnalysisEpisodeMember.objects.filter(Q(episode_id__in=episode_ids)|Q(raw_event__site_id__in=site_pks)).delete()[0]
        deleted_episodes=AnalysisEpisode.objects.filter(pk__in=episode_ids).delete()[0]
        count,details=Site.objects.filter(pk__in=site_pks).delete()
    return count,details,deleted_members,deleted_episodes


def rebuild_retained_scope(result,progress=print):
    root=Path(settings.BASE_DIR)
    # Preserve previous labels as history, then recompute DBSCAN on the retained
    # inventory. Reset aliases first so old/new real labels cannot conflict.
    if result.get('scope')=='tdd_cells':
        # The site inventory and coordinates are unchanged. Reuse the original
        # spatial clusters and refresh cell-dependent neighborhood features.
        result['clustering']={'method':'original site-only assignments retained; inventory and coordinates unchanged'}
    else:
        Site.objects.filter(original_cluster_id__isnull=True).update(original_cluster_id=F('cluster_id'))
        Site.objects.update(cluster_id=-1,is_cluster_edge=False,canonical_site_id='')
        from services.models.preprocessing.cellplan_loader import CellPlanLoader
        clustering=CellPlanLoader(run_clustering=True).run_clustering_pipeline()
        result['clustering']={key:int(clustering[key]) for key in ('n_clusters','n_noise','out_of_bounds_count')}
    from services.models.preprocessing.network_events import detect_and_save_events
    detection=detect_and_save_events()
    result['detection']={key:value for key,value in detection.items() if isinstance(value,(int,float,str,bool,type(None)))}
    from services.models.features.pipeline import FeatureExtractionPipeline
    def callback(data):
        if data.get('processed',0)==0 or (data.get('processed',0)+1)%100==0 or data.get('processed')==data.get('total'):
            progress(f"Feature refresh: {data}")
    result['features']=FeatureExtractionPipeline().extract_and_save_features(skip_existing=False,progress_callback=callback,reuse_static=result.get('scope')=='tdd_cells')
    from django.core.management import call_command
    for command in ('freeze_similarity_reference','build_pair_cases'):
        progress(f'Rebuilding {command}');call_command(command)
    progress('Rebuilding episode recipient cases');call_command('build_pair_cases',episodes=True)
    progress('Rebuilding gain recipient cases');call_command('build_gain_cases')
    result['rebuild_status']='complete'
    result['remaining_sites']=Site.objects.count();result['remaining_cells']=Cell.objects.count();result['remaining_kpis']=MAOSMonthlyKPI.objects.count()
    result['remaining_physical_sites']=Site.objects.values('canonical_site_id').distinct().count()
    result['remaining_cband_cells']=sum((tdd_cell(c.band,c.frequency,c.technology) if result.get('scope')=='tdd_cells' else plan_cband(c.band,c.frequency,c.technology)) for c in Cell.objects.all())
    if result['remaining_cband_cells'] or Site.objects.filter(site_id__in=scope_policy()['excluded_site_ids']).exists():
        raise ValueError('Post-cleanup C-band exclusion validation failed.')
    backup=root/result['backup'];filename='tdd_cell_removal.json' if result.get('scope')=='tdd_cells' else 'removal.json'
    (backup/filename).write_text(json.dumps(result,indent=2),encoding='utf-8')
    output=root/'docs/verification/data_scope'/filename;output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text(json.dumps(result,indent=2),encoding='utf-8')
    return result
